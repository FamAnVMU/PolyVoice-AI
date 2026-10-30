# -*- coding: utf-8 -*-
"""
Script to generate the comprehensive system documentation for PolyVoice AI in DOCX format.
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_callout_box(paragraph, border_color="6C5CE7", bg_color="F0EEFC"):
    pPr = paragraph._p.get_or_add_pPr()
    # Left border
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="15" w:color="{border_color}"/></w:pBdr>')
    pPr.append(pBdr)
    # Shading
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    pPr.append(shd)

def create_document():
    doc = docx.Document()

    # Page Margins: 0.8 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles & Fonts
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x2D, 0x34, 0x36)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(6)

    # Palette
    COLOR_PRIMARY = RGBColor(0x48, 0x34, 0xD4)    # Deep Indigo
    COLOR_SECONDARY = RGBColor(0x00, 0x94, 0x85)  # Teal
    COLOR_DARK = RGBColor(0x1F, 0x24, 0x2E)       # Dark Slate
    COLOR_MUTED = RGBColor(0x63, 0x6E, 0x72)      # Gray

    # -------------------------------------------------------------
    # COVER / HEADER TITLE
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("TÀI LIỆU MÔ TẢ KIẾN TRÚC VÀ ĐẶC TẢ HỆ THỐNG")
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("ỨNG DỤNG POLYVOICE AI (PHIÊN BẢN 1.1.0)")
    r_sub.font.size = Pt(15)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    desc_p = doc.add_paragraph()
    desc_p.paragraph_format.space_after = Pt(20)
    desc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_desc = desc_p.add_run("Nền Tảng Giao Tiếp & Đàm Thoại Thời Gian Thực Đa Nhân Cách AI\nKiến Trúc Client-Side Zero-Server-Cost (BYOK) • Multi-LLM • Multi-Voice Engine")
    r_desc.font.size = Pt(11)
    r_desc.font.italic = True
    r_desc.font.color.rgb = COLOR_MUTED

    # Meta Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Dự án:", "PolyVoice AI (Flutter Multiplatform)"),
        ("Phiên bản:", "1.1.0+8 (Major Architecture Release)"),
        ("Mã gói (Application ID):", "com.polyvoice.polyvoice_ai"),
        ("Phân loại tài liệu:", "Technical Architecture & System Specification"),
        ("Ngày phát hành:", "30/10/2026")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.5)
        cell_k.paragraphs[0].text = k
        cell_k.paragraphs[0].runs[0].font.bold = True
        cell_k.paragraphs[0].runs[0].font.color.rgb = COLOR_DARK
        cell_v.paragraphs[0].text = v
        set_cell_background(cell_k, "F1F2F6")
        set_cell_background(cell_v, "FAFAFA")
        set_cell_margins(cell_k, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_v, top=60, bottom=60, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    # Divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(20)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="6C5CE7"/></w:pBdr>')
    div_p._p.get_or_add_pPr().append(pBdr)

    # -------------------------------------------------------------
    # HELPER FUNCTIONS FOR SECTIONS
    # -------------------------------------------------------------
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(22)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY
        return h

    def add_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_DARK
        return h

    def add_callout(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Inches(0.15)
        set_callout_box(p, border_color="6C5CE7", bg_color="F4F2FF")
        if bold_prefix:
            r_pre = p.add_run(bold_prefix + " ")
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_PRIMARY
        r_txt = p.add_run(text)
        r_txt.font.size = Pt(10.5)
        return p

    def format_table(table, col_widths, headers, data):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Header Row
        hdr_row = table.rows[0]
        for idx, title in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.width = Inches(col_widths[idx])
            cell.paragraphs[0].text = title
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cell.paragraphs[0].runs[0]
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_cell_background(cell, "4834D4")
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)

        # Data Rows
        for row_idx, row_data in enumerate(data):
            row = table.rows[row_idx + 1]
            bg_color = "F8F9FE" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_value in enumerate(row_data):
                cell = row.cells[col_idx]
                cell.width = Inches(col_widths[col_idx])
                cell.paragraphs[0].text = str(cell_value)
                run = cell.paragraphs[0].runs[0]
                run.font.size = Pt(9.5)
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # -------------------------------------------------------------
    # CHƯƠNG 1: TỔNG QUAN HỆ THỐNG
    # -------------------------------------------------------------
    add_h1("CHƯƠNG 1: TỔNG QUAN DỰ ÁN & TRIẾT LÝ THIẾT KẾ")

    p = doc.add_paragraph("PolyVoice AI là ứng dụng trợ lý thông minh đa nhân cách thế hệ mới, kết hợp khả năng giao tiếp văn bản (Text Chat) và hội thoại giọng nói hai chiều trực tiếp (Real-Time Continuous Voice Chat). Ứng dụng được xây dựng trên nền tảng Flutter 3.x, vận hành theo mô hình kiến trúc không phụ thuộc máy chủ trung gian (Zero-Server-Cost), mang lại quyền riêng tư và khả năng tùy biến vượt trội cho người sử dụng.")

    add_h2("1.1. Các Giá Trị Cốt Lõi Của Hệ Thống")
    doc.add_paragraph("Hệ thống PolyVoice AI được kiến tạo dựa trên 4 trụ cột công nghệ then chốt:")

    b1 = doc.add_paragraph(style='List Bullet')
    b1.add_run("Mô hình Client-Side & BYOK (Bring Your Own Key): ").bold = True
    b1.add_run("Ứng dụng kết nối trực tiếp từ thiết bị người dùng đến các API nhà cung cấp AI hàng đầu thế giới (Google Gemini, Groq, OpenRouter, ElevenLabs, Microsoft Edge-TTS). Không có bất kỳ máy chủ trung gian nào lưu trữ lén tin nhắn, âm thanh hoặc API Key.")

    b2 = doc.add_paragraph(style='List Bullet')
    b2.add_run("Nền tảng Mạng Xã Hội Đa Nhân Cách AI (Multi-Character Social Platform): ").bold = True
    b2.add_run("Chuyển dịch hoàn toàn từ trải nghiệm bot đơn điệu sang hệ sinh thái gồm nhiều nhân vật AI độc lập thuộc 5 danh mục: Bạn bè, Học tập, Chuyên gia, Giải trí & Nhập vai, Tâm sự & Chữa lành. Mỗi nhân vật sở hữu tên tuổi, avatar, tiểu sử, giọng nói và bộ nhớ tách biệt tuyệt đối.")

    b3 = doc.add_paragraph(style='List Bullet')
    b3.add_run("Động cơ Đàm Thoại Thời Gian Thực (Low-Latency Voice Engine): ").bold = True
    b3.add_run("Hệ thống Voice Pipeline Coordinator kết hợp bộ nhận diện giọng nói STT, bộ phân đoạn câu thông minh (Sentence Chunker) và hàng đợi phát âm thanh TtsQueueManager cho phép phát âm thanh nối tiếp mượt mà ngay trong lúc AI đang tiếp tục truyền tải luồng văn bản (Streaming).")

    b4 = doc.add_paragraph(style='List Bullet')
    b4.add_run("Cá Nhân Hóa Toàn Diện & Bảo Toàn Dữ Liệu: ").bold = True
    b4.add_run("Hệ thống tự động tinh chỉnh ngữ điệu và nội dung câu trả lời theo từng nhóm độ tuổi (Thiếu nhi, Học sinh/SV, Trưởng thành, Người lớn tuổi); hỗ trợ sao lưu và phục hồi dữ liệu 100% bằng chuẩn JSON, đồng thời cam kết tương thích cài đè cập nhật liên tục qua chứng chỉ ký số vĩnh viễn.")

    add_callout(
        "Phiên bản 1.1.0 đánh dấu bước tiến lớn của PolyVoice AI từ một chatbot truyền thống thành một nền tảng xã hội AI đa nhân cách hoàn chỉnh, giải quyết triệt để lỗi xung đột cài đè Android, lỗi phân mảnh câu trả lời streaming và lỗi rò rỉ ngữ cảnh.",
        "Điểm mốc phiên bản 1.1.0:"
    )

    # -------------------------------------------------------------
    # CHƯƠNG 2: KIẾN TRÚC PHẦN MỀM TỔNG THỂ
    # -------------------------------------------------------------
    add_h1("CHƯƠNG 2: KIẾN TRÚC PHẦN MỀM TỔNG THỂ (SOFTWARE ARCHITECTURE)")

    doc.add_paragraph("PolyVoice AI áp dụng triệt để mô hình Clean Architecture kết hợp State Management BLoC (Business Logic Component). Kiến trúc được module hóa theo từng tầng tính năng rõ rệt, tách rời giữa Presentation, Domain/Logic và Data/Infrastructure.")

    add_h2("2.1. Sơ Đồ Phân Tầng Kiến Trúc Hệ Thống")

    arch_table = doc.add_table(rows=6, cols=3)
    arch_widths = [1.6, 2.0, 3.1]
    arch_headers = ["Tầng Kiến Trúc", "Thành Phần Chính", "Trách Nhiệm & Chức Năng"]
    arch_data = [
        ("Presentation Layer\n(Giao diện & UI)", "MainNavigationScreen\nHomeTab, DiscoverTab\nMessagesTab, ProfileTab\nChatScreen, ContinuousVoice", "Hiển thị giao diện 4 tab xã hội, quản lý tương tác cảm ứng, micro, hoạt họa sóng âm, danh sách tin nhắn và hồ sơ cá nhân."),
        ("BLoC State Layer\n(Quản lý trạng thái)", "ChatBloc (ChatState/Event)\nSettingsBloc (SettingsState/Event)", "Điều phối luồng trạng thái bất đồng bộ, tiếp nhận sự kiện người dùng, tương tác với kho lưu trữ và điều khiển Voice Coordinator."),
        ("Pipeline Coordinator\n(Điều phối đàm thoại)", "VoicePipelineCoordinator\nSentenceChunker\nStreamResponseSanitizer\nVietnameseTextNormalizer", "Điều phối chuỗi STT -> LLM Stream -> Sanitizer -> Chunking -> TTS Queue. Đảm bảo âm thanh phát ra có độ trễ cực thấp và chuẩn hóa ngữ âm."),
        ("Services & Providers\n(Dịch vụ & AI Engine)", "LlmServiceFactory (Gemini, Groq, OpenRouter)\nEdgeTtsService, ElevenLabs\nTtsService, SttService", "Đóng gói toàn bộ logic kết nối API bên thứ ba. Hỗ trợ truyền dữ liệu SSE (Server-Sent Events) đa dòng và chuyển đổi văn bản sang âm thanh."),
        ("Storage & Security\n(Lưu trữ & Bảo mật)", "ChatLocalStorageService\nSecureStorageService\nVoiceStorageService\nRelease Keystore", "Lưu trữ phiên chat (Schema v2 Migration), mã hóa API Key qua Hardware Keystore/Keychain, xuất/nhập tệp sao lưu JSON.")
    ]
    format_table(arch_table, arch_widths, arch_headers, arch_data)

    add_h2("2.2. Cấu Trúc Thư Mục Mã Nguồn (Directory Structure)")
    doc.add_paragraph("Mã nguồn dự án được tổ chức khoa học theo mô hình Feature-First trong thư mục `lib/`:")

    dir_box = doc.add_paragraph()
    set_callout_box(dir_box, border_color="009485", bg_color="F5FCFB")
    dir_box.add_run(
        "lib/\n"
        "├── core/                 # Cấu hình cốt lõi (Theme, AppColors, Errors, Logger, Utils)\n"
        "├── features/             # Các phân hệ tính năng người dùng\n"
        "│   ├── navigation/       # MainNavigationScreen điều hướng 4 tab\n"
        "│   ├── home/             # HomeTab (Chào hỏi, Banner, Nhân vật nổi bật, Chat gần đây)\n"
        "│   ├── discover/         # DiscoverTab (Tìm kiếm, Bộ lọc chuyên mục, Bộ lọc độ tuổi)\n"
        "│   ├── chat/             # MessagesTab, ChatScreen, ChatBloc, Models, Widgets\n"
        "│   ├── character/        # AICharacterModel, CharacterLibrary (16 nhân vật), ProfileDialog\n"
        "│   ├── profile/          # UserProfileModel, ProfileTab (Hồ sơ, Tuổi, Sao lưu/Khôi phục)\n"
        "│   ├── settings/         # SettingsBloc, ApiKeysView, VoiceShaperView, PersonaEditorView\n"
        "│   └── voice/            # ContinuousVoiceScreen, VoiceModel, PresetVoices\n"
        "└── services/             # Các tầng dịch vụ chuyên sâu\n"
        "    ├── pipeline/         # VoicePipelineCoordinator (Orchestrator điều phối STT/TTS)\n"
        "    ├── llm/              # LlmServiceFactory, GeminiService, GroqService, OpenRouterService\n"
        "    ├── voice/            # SttService, TtsService, TtsQueueManager, EdgeTts, ElevenLabs\n"
        "    └── storage/          # ChatLocalStorageService, SecureStorageService"
    )
    dir_box.runs[0].font.name = 'Consolas'
    dir_box.runs[0].font.size = Pt(9)

    # -------------------------------------------------------------
    # CHƯƠNG 3: HỆ THỐNG ĐA NHÂN CÁCH VÀ BỘ NHỚ ĐỘC LẬP
    # -------------------------------------------------------------
    add_h1("CHƯƠNG 3: HỆ THỐNG ĐA NHÂN CÁCH AI VÀ BỘ NHỚ ĐỘC LẬP")

    doc.add_paragraph("Điểm nâng cấp mang tính bản lề của PolyVoice AI là việc biến mỗi cuộc trò chuyện thành một thực thể nhân vật AI độc lập với hệ thống tính cách, phong cách đối thoại và kho ký ức hoàn toàn riêng biệt.")

    add_h2("3.1. Thư Viện 16 Nhân Vật AI Chuyên Biệt")
    doc.add_paragraph("Thư viện CharacterLibrary cung cấp 16 nhân vật được xây dựng công phu, phân loại thành 5 danh mục ứng dụng:")

    char_table = doc.add_table(rows=17, cols=5)
    char_widths = [1.2, 0.5, 1.3, 1.4, 2.3]
    char_headers = ["Tên Nhân Vật", "Icon", "Danh Mục", "Giọng Đọc Mặc Định", "Đặc Điểm & Phong Thái Đối Thoại"]
    char_data = [
        ("Khánh Linh", "🌸", "Bạn bè", "preset_khanh_linh", "Dịu dàng, ấm áp, biết lắng nghe, xưng Bạn - Linh. Luôn quan tâm và chia sẻ tâm sự."),
        ("Gia Huy", "⚡", "Bạn bè", "preset_minh_an", "Hài hước, dí dỏm, tràn đầy năng lượng Gen Z, xưng Bạn - Huy/Tui."),
        ("Thanh Trúc", "🍃", "Bạn bè", "preset_ngoc_anh", "Tinh tế, sâu lắng, thích nghệ thuật và những cuộc trò chuyện chiêm nghiệm cuộc sống."),
        ("Bảo Long", "🏔️", "Bạn bè", "preset_anh_khoi", "Điềm đạm, chững chạc, góc nhìn khách quan, là điểm tựa đáng tin cậy."),
        ("Thầy Minh Quân", "📚", "Học tập", "preset_minh_an", "Sư phạm mẫu mực, kiên nhẫn giảng giải khoa học tự nhiên, toán học và tư duy logic."),
        ("Cô Hoàng Yến", "🖋️", "Học tập", "preset_ngoc_anh", "Truyền cảm hứng, giàu vốn từ, chuyên sâu về văn học, lịch sử và kỹ năng viết."),
        ("TS. Tuấn Kiệt", "🔬", "Học tập", "preset_anh_khoi", "Học thuật chuẩn xác, tư duy nghiên cứu, phân tích đa chiều và phương pháp luận."),
        ("Alex Bilingual", "🌍", "Học tập", "preset_khanh_linh", "Song ngữ Anh - Việt linh hoạt, sửa lỗi ngữ pháp, luyện phản xạ giao tiếp quốc tế."),
        ("Dev Senior Vũ", "💻", "Chuyên gia", "preset_anh_khoi", "Thực chiến IT, clean code, giải quyết bug hóc búa, thiết kế kiến trúc phần mềm."),
        ("Văn Hưng", "📊", "Chuyên gia", "preset_minh_an", "Quản lý tài chính cá nhân, đầu tư thực tế, quản trị dòng tiền và phân tích kinh doanh."),
        ("Mai Phương", "⚖️", "Chuyên gia", "preset_ngoc_anh", "Cẩn trọng, chặt chẽ, tư vấn pháp lý dân sự, thủ tục hợp đồng và quyền lợi người dân."),
        ("Coach Hải Đăng", "🏃", "Chuyên gia", "preset_anh_khoi", "Kỷ luật thép, tràn đầy năng lượng, thúc đẩy rèn luyện thể chất và đạt mục tiêu."),
        ("Thảo Nhi", "🎭", "Giải trí", "preset_khanh_linh", "Hoạt ngôn, vui tươi, đố vui thông minh, mini-game tương tác và truyện cười hài hước."),
        ("Hiệp sĩ Arthur", "🛡️", "Nhập vai", "preset_anh_khoi", "Quả cảm, danh dự hiệp sĩ thời trung cổ, phiêu lưu kỳ ảo và đồng hành trung kiên."),
        ("MC Quốc Huy", "🎙️", "Giải trí", "preset_minh_an", "Hoạt náo viên talkshow, đối thoại sắc sảo, tranh biện và dẫn dắt chủ đề hấp dẫn."),
        ("Tâm An", "🕊️", "Chữa lành", "preset_ngoc_anh", "Tĩnh tại, trắc ẩn, xoa dịu stress, hướng dẫn thiền định và lắng nghe không phán xét.")
    ]
    format_table(char_table, char_widths, char_headers, char_data)

    add_h2("3.2. Cơ Chế Tiêm Prompt Động 3 Tầng (3-Layer Dynamic Prompt Injection)")
    doc.add_paragraph("Mỗi khi người dùng gửi một tin nhắn văn bản hoặc giọng nói, ChatBloc tổng hợp một System Prompt động hoàn chỉnh trước khi gửi đến mô hình LLM:")

    add_callout(
        "TẦNG 1: System Prompt Cốt Lõi của Nhân Vật (Character System Prompt)\n"
        "Định hình tên, đại từ xưng hô, tính cách, giới hạn đạo đức và phong thái chuyên môn của nhân vật.\n\n"
        "TẦNG 2: Chỉ Dẫn Cá Nhân Hóa Người Dùng (User Profile & Age-Group Addon)\n"
        "Tiêm hướng dẫn đối tượng (Thiếu nhi: từ ngữ trong sáng, an toàn; Học sinh: khuyến khích tư duy; Người lớn: luận điểm logic sâu sắc; Người cao tuổi: kính trọng, từ tốn), phong cách trả lời (Chi tiết / Súc tích) và đại từ xưng hô mong muốn.\n\n"
        "TẦNG 3: Ký Ức Riêng Của Cuộc Trò Chuyện (Session Memory Notes)\n"
        "Tiêm các ghi chú sự kiện riêng biệt đã lưu trong cuộc hội thoại đó (ví dụ: 'Người dùng đang chuẩn bị bảo vệ đồ án tốt nghiệp').",
        "Cấu trúc System Prompt Động 3 Tầng:"
    )

    add_h2("3.3. Cô Lập Ngữ Cảnh Tuyệt Đối (Context & Memory Isolation)")
    doc.add_paragraph("Trong mô hình ChatSession, mỗi phiên chat được cấp một UUID v4 độc lập cùng mảng `memoryNotes` riêng. Lịch sử trò chuyện và ghi nhớ của phiên này hoàn toàn không thể rò rỉ sang phiên khác. Người dùng có thể nhấn vào Header của ChatScreen để chủ động xem lại, bổ sung hoặc xóa bỏ bất kỳ mục ký ức nào của nhân vật đối thoại.")

    # -------------------------------------------------------------
    # CHƯƠNG 4: THIẾT KẾ GIAO DIỆN VÀ TRẢI NGHIỆM NGƯỜI DÙNG
    # -------------------------------------------------------------
    add_h1("CHƯƠNG 4: THIẾT KẾ GIAO DIỆN & TRẢI NGHIỆM NGƯỜI DÙNG (UI/UX)")

    doc.add_paragraph("Giao diện PolyVoice AI được thiết kế theo xu hướng mạng xã hội AI hiện đại với hệ màu Indigo (#6C5CE7) kết hợp Neon Cyan (#00CEC9), hỗ trợ chuyển đổi mượt mà giữa Giao diện Tối (Dark Mode) và Giao diện Sáng (Light Mode).")

    add_h2("4.1. Cấu Trúc Điều Hướng 4 Tab (MainNavigationScreen)")
    doc.add_paragraph("Người dùng dễ dàng di chuyển giữa các phân hệ chức năng thông qua thanh điều hướng đáy cố định:")

    b_t1 = doc.add_paragraph(style='List Bullet')
    b_t1.add_run("Tab 1 - Trang Chủ (HomeTab): ").bold = True
    b_t1.add_run("Trung tâm điều khiển cá nhân hóa với lời chào theo giờ, hiển thị nhóm tuổi, nút gọi thoại khẩn cấp, biểu ngữ khám phá, băng chuyền lướt ngang các nhân vật nổi bật và danh sách các cuộc trò chuyện gần nhất.")

    b_t2 = doc.add_paragraph(style='List Bullet')
    b_t2.add_run("Tab 2 - Khám Phá (DiscoverTab): ").bold = True
    b_t2.add_run("Thư viện tìm kiếm và phân loại nhân vật. Cung cấp bộ lọc theo 5 nhóm danh mục, bộ lọc độ tuổi phù hợp và lưới thẻ nhân vật 2 cột với đánh giá sao và nút nhắn tin nhanh.")

    b_t3 = doc.add_paragraph(style='List Bullet')
    b_t3.add_run("Tab 3 - Tin Nhắn (MessagesTab): ").bold = True
    b_t3.add_run("Hộp thư đa phiên với 3 phân loại: Tất Cả, Đã Ghim (Pinned) và Lưu Trữ (Archived). Hỗ trợ tìm kiếm hội thoại, thao tác ghim/lưu trữ/xóa và nút FloatingActionButton khởi tạo cuộc trò chuyện mới.")

    b_t4 = doc.add_paragraph(style='List Bullet')
    b_t4.add_run("Tab 4 - Cá Nhân (ProfileTab): ").bold = True
    b_t4.add_run("Quản lý hồ sơ người dùng, lựa chọn emoji đại diện, tùy chỉnh nhóm lứa tuổi, phong cách câu trả lời, thiết lập API Keys, phòng thu Voice Studio, công cụ sao lưu/khôi phục JSON và kiểm tra chữ ký Release Keystore.")

    add_h2("4.2. Màn Hình Hội Thoại (ChatScreen) & Đàm Thoại Trực Tiếp (ContinuousVoiceScreen)")
    doc.add_paragraph("Màn hình ChatScreen hiển thị tên và avatar của nhân vật đang mở, huy hiệu BYOK trực quan, hỗ trợ tự động cuộn thông minh (Auto-Scroll Lock khi người dùng lướt lên xem tin cũ), hiển thị Markdown đẹp mắt và nút phát giọng đọc TTS cho từng tin nhắn. Màn hình ContinuousVoiceScreen cung cấp giao diện Voice Orb sống động với 4 trạng thái phản hồi quang học (Chờ, Lắng nghe, Suy nghĩ, Phát âm).")

    # -------------------------------------------------------------
    # CHƯƠNG 5: ĐỘNG CƠ XỬ LÝ GIỌNG NÓI THỜI GIAN THỰC
    # -------------------------------------------------------------
    add_h1("CHƯƠNG 5: ĐỘNG CƠ XỬ LÝ GIỌNG NÓI THỜI GIAN THỰC (VOICE PIPELINE)")

    doc.add_paragraph("Trọng tâm kỹ thuật giúp PolyVoice AI vượt trội so với các ứng dụng thông thường là đường ống xử lý giọng nói tích hợp liên tục (End-to-End Voice Orchestration Pipeline) do VoicePipelineCoordinator điều phối.")

    add_h2("5.1. Chuỗi Xử Lý Âm Thanh Bất Đồng Bộ (Pipeline Workflow)")
    doc.add_paragraph("Chu kỳ xử lý một lượt đàm thoại giọng nói tuân thủ theo các bước sau:")

    pipe_table = doc.add_table(rows=7, cols=3)
    pipe_widths = [1.2, 1.8, 3.7]
    pipe_headers = ["Giai Đoạn", "Module Phụ Trách", "Nhiệm Vụ Kỹ Thuật"]
    pipe_data = [
        ("1. Thu âm & STT", "SttService\n(Speech-To-Text)", "Thu nhận tín hiệu từ Micro, kích hoạt VAD (Voice Activity Detection), hiển thị Live Transcript Preview thời gian thực."),
        ("2. Truy vấn LLM", "LlmServiceFactory\n(Streaming SSE)", "Gửi prompt và ngữ cảnh hội thoại đến LLM được chọn, tiếp nhận dòng dữ liệu Server-Sent Events đa dòng không gián đoạn."),
        ("3. Lọc phản hồi", "StreamResponseSanitizer", "Lọc bỏ tức thì các thẻ suy nghĩ trung gian (<think>...</think>), các nhãn 'Draft response:' hoặc ký tự rác của AI."),
        ("4. Phân đoạn câu", "SentenceChunker", "Phát hiện dấu câu tiếng Việt (. ? ! : \\n) để cắt chuỗi thành các phân đoạn ngữ nghĩa hoàn chỉnh có độ dài tối ưu (40 - 150 ký tự)."),
        ("5. Chuẩn hóa ngữ âm", "VietnameseTextNormalizer", "Chuyển đổi số, tỷ lệ phần trăm, ngày tháng, tiền tệ (VNĐ, USD), ký tự toán học thành chữ viết trước khi nạp vào TTS."),
        ("6. Phát âm thanh", "TtsQueueManager\n(EdgeTTS / ElevenLabs)", "Xếp các phân đoạn vào hàng đợi âm thanh. Khi một đoạn đang phát, đoạn tiếp theo đã được tải sẵn (pre-buffering). Cho phép ngắt ngay khi người dùng nói đè.")
    ]
    format_table(pipe_table, pipe_widths, pipe_headers, pipe_data)

    add_h2("5.2. Các Động Cơ Giọng Nói (Voice Engines) Hỗ Trợ")
    doc.add_paragraph("Hệ thống cung cấp sự linh hoạt tối đa về chất lượng và chi phí âm thanh:")
    doc.add_paragraph("• Microsoft Edge-TTS: Giọng đọc nơ-ron miễn phí, tự nhiên, tốc độ cao, hỗ trợ đầy đủ tiếng Việt (Nam Minh, Nữ Hoài My).\n• ElevenLabs AI: Giọng siêu thực điện ảnh, hỗ trợ tinh chỉnh độ ổn định (Stability), cảm xúc (Emotion) và phong cách (Style).\n• Native System TTS: Động cơ giọng đọc ngoại tuyến của hệ điều hành, đảm bảo ứng dụng luôn có tiếng nói dù mạng chập chờn.")

    # -------------------------------------------------------------
    # CHƯƠNG 6: TÍCH HỢP MÔ HÌNH NGÔN NGỮ LỚN (MULTI-LLM)
    # -------------------------------------------------------------
    add_h1("CHƯƠNG 6: TÍCH HỢP MÔ HÌNH NGÔN NGỮ LỚN (MULTI-LLM BYOK)")

    doc.add_paragraph("PolyVoice AI cung cấp khả năng tích hợp đa mô hình ngôn ngữ thông qua giao diện LlmProvider đồng nhất, người dùng tự do lựa chọn nhà cung cấp AI phù hợp với nhu cầu và hạn mức:")

    llm_table = doc.add_table(rows=4, cols=4)
    llm_widths = [1.5, 1.8, 1.8, 1.6]
    llm_headers = ["Nhà Cung Cấp", "Mô Hình Đề Xuất", "Đặc Điểm Kỹ Thuật", "Ưu Thế Sử Dụng"]
    llm_data = [
        ("Google Gemini", "gemini-2.5-flash\ngemini-1.5-pro", "Hỗ trợ ngữ cảnh siêu dài, tốc độ xử lý nhanh, nhận diện tiếng Việt xuất sắc.", "Miễn phí định mức cao, tối ưu cho đàm thoại đa chủ đề."),
        ("Groq Cloud", "llama-3.3-70b-versatile\nmixtral-8x7b-32768", "Tốc độ tạo token cực hạn (LPU Inference Engine), độ trễ TTFT cực thấp.", "Lý tưởng nhất cho chế độ Continuous Voice Chat thời gian thực."),
        ("OpenRouter", "claude-3-5-sonnet\ngpt-4o / deepseek-chat", "Cổng kết nối hàng trăm mô hình mã nguồn mở và thương mại hàng đầu thế giới.", "Lựa chọn tối thượng cho người dùng nâng cao muốn thử nghiệm đa dạng.")
    ]
    format_table(llm_table, llm_widths, llm_headers, llm_data)

    add_callout(
        "Nhờ cơ chế Multi-line SSE Buffer mới trong gemini_service.dart, các câu trả lời dài của AI không còn bị cắt cụt sau dòng đầu tiên, bảo đảm chất lượng phân tích chuyên sâu toàn vẹn.",
        "Khắc phục sự cố truyền tải SSE:"
    )

    # -------------------------------------------------------------
    # CHƯƠNG 7: BẢO MẬT, QUẢN LÝ PHIÊN BẢN VÀ DỮ LIỆU
    # -------------------------------------------------------------
    add_h1("CHƯƠNG 7: BẢO MẬT, QUẢN LÝ PHIÊN BẢN & BẢO TOÀN DỮ LIỆU")

    add_h2("7.1. Khắc Phục Triệt Để Xung Đột Chữ Ký Cài Đè (APK Signing Config)")
    doc.add_paragraph("Trước phiên bản 1.1.0, các bản build từ GitHub Actions gặp lỗi nghiêm trọng INSTALL_FAILED_UPDATE_INCOMPATIBLE do mỗi runner tự sinh keystore debug ngẫu nhiên. Giải pháp đã thực hiện:")
    doc.add_paragraph("• Sinh tệp `release-keystore.jks` chuẩn RSA 2048-bit với thời hạn 10.000 ngày (~27 năm) và ghim trực tiếp vào repository.\n• Cấu hình build.gradle đồng bộ chữ ký này cho cả hai loại build (Release & Debug).\n• Giữ nguyên định danh gói `com.polyvoice.polyvoice_ai` xuyên suốt, bảo đảm người dùng cập nhật phiên bản mới 100% không bị xung đột.")

    add_h2("7.2. Bảo Mật API Key Cục Bộ (Hardware-Backed Keystore)")
    doc.add_paragraph("Toàn bộ API Key người dùng nhập vào được bảo vệ bởi SecureStorageService. Dữ liệu được mã hóa bằng thuật toán AES phần cứng, lưu trong Android Keystore hoặc iOS Keychain, ngăn chặn triệt để nguy cơ trích xuất từ các ứng dụng độc hại khác.")

    add_h2("7.3. Nâng Cấp Cấu Trúc Dữ Liệu Tự Động (Database Migration v2)")
    doc.add_paragraph("Khi ứng dụng nâng cấp từ bản cũ lên bản 1.1.0, hàm `_migrateFromV1ToV2` trong ChatLocalStorageService tự động bổ sung các trường dữ liệu nhân cách và ghim/lưu trữ cho các phiên chat cũ mà không làm mất bất kỳ một tin nhắn lịch sử nào.")

    add_h2("7.4. Công Cụ Sao Lưu Và Khôi Phục Dữ Liệu JSON Độc Lập")
    doc.add_paragraph("Người dùng có toàn quyền sở hữu dữ liệu của mình: chỉ với một nút bấm trong tab Cá nhân, toàn bộ hồ sơ, các phiên trò chuyện, các mục ký ức nhân cách được đóng gói thành chuỗi JSON tiêu chuẩn để sao lưu hoặc chuyển giao sang thiết bị mới mà không cần tài khoản đám mây.")

    # -------------------------------------------------------------
    # CHƯƠNG 8: KẾ HOẠCH PHÁT TRIỂN TƯƠNG LAI (ROADMAP)
    # -------------------------------------------------------------
    add_h1("CHƯƠNG 8: KẾ HOẠCH PHÁT TRIỂN TƯƠNG LAI (ROADMAP)")

    road_table = doc.add_table(rows=5, cols=3)
    road_widths = [1.2, 2.0, 3.5]
    road_headers = ["Giai Đoạn", "Tính Năng Dự Kiến", "Mô Tả Kỹ Thuật"]
    road_data = [
        ("Giai đoạn 1\n(Hiện tại v1.1.0)", "Nền tảng Mạng Xã Hội AI\n& Release Keystore", "Hoàn thiện 4 tab xã hội, 16 nhân vật AI, bộ nhớ riêng biệt, BYOK đa nhà cung cấp, bảo toàn dữ liệu cập nhật cài đè."),
        ("Giai đoạn 2\n(v1.2.0)", "Live Bidirectional WebSockets\n(Gemini 2.0 Live API)", "Nâng cấp giao thức hội thoại giọng nói sang WebSocket hai chiều thời gian thực (Native Audio In -> Native Audio Out), độ trễ < 300ms."),
        ("Giai đoạn 3\n(v1.3.0)", "Local RAG & Knowledge Base\n(Tài liệu cục bộ)", "Tích hợp cơ sở dữ liệu véc-tơ cục bộ (ObjectBox / Isar), cho phép nhân vật AI đọc hiểu tệp PDF, TXT do người dùng tải lên."),
        ("Giai đoạn 4\n(v2.0.0)", "Custom Character Creator\n(Tự tạo nhân vật)", "Cho phép người dùng tự thiết kế nhân vật AI mới: tải ảnh đại diện, viết System Prompt, chọn giọng đọc và chia sẻ qua mã QR/JSON.")
    ]
    format_table(road_table, road_widths, road_headers, road_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(24)

    # Footer note
    end_p = doc.add_paragraph()
    end_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = end_p.add_run("--- HẾT TÀI LIỆU MÔ TẢ HỆ THỐNG POLYVOICE AI ---")
    r_end.font.bold = True
    r_end.font.color.rgb = COLOR_MUTED
    r_end.font.size = Pt(10)

    # Save
    out_path = r"c:\Users\phama\Downloads\assistant app\Tai_Lieu_Mo_Ta_He_Thong_PolyVoice_AI.docx"
    doc.save(out_path)
    print(f"Document created successfully at: {out_path}")

if __name__ == "__main__":
    create_document()
